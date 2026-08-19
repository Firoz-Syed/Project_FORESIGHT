from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path

output_path = Path("reports/Project_FORESIGHT_Final_Presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slides = [
    (
        "Project FORESIGHT",
        "Retail Demand Forecasting & Inventory Risk Analysis using Machine Learning"
    ),
    (
        "Problem Statement",
        "Retail businesses face demand uncertainty, stockout risk, overstocking, "
        "and inefficient replenishment planning."
    ),
    (
        "Project Objectives",
        "• Forecast retail product demand\n"
        "• Identify inventory risk\n"
        "• Support replenishment decisions\n"
        "• Generate business insights\n"
        "• Provide interactive analytics"
    ),
    (
        "Dataset & Data Engineering",
        "73,000 final records\n"
        "200 unique SKUs\n"
        "2024 analysis period\n"
        "24 final dataset columns\n"
        "Source data: sales, inventory, SKU master, and calendar datasets"
    ),
    (
        "EDA & Feature Engineering",
        "Exploratory analysis identified sales, revenue, inventory, category, "
        "seasonal, calendar, and promotional patterns.\n\n"
        "Feature engineering produced model-ready temporal and inventory features."
    ),
    (
        "Machine Learning Model",
        "Final model: Random Forest Regressor\n"
        "50 estimators\n"
        "15 input features\n"
        "Time-based 80/20 train-test split\n"
        "Training rows: 58,400\n"
        "Testing rows: 14,600"
    ),
    (
        "Model Evaluation",
        "Final Test Metrics\n"
        "MAE: 3.2449\n"
        "RMSE: 4.0551\n"
        "R²: 0.0946\n\n"
        "Final model: models/final_random_forest.pkl"
    ),
    (
        "Inventory Risk Scoring",
        "73,000 records evaluated\n"
        "High Risk: 63,400 (86.85%)\n"
        "Medium Risk: 51 (0.07%)\n"
        "Low Risk: 9,549 (13.08%)\n\n"
        "Critical risk output fields contain no missing values."
    ),
    (
        "Business Insights",
        "High-risk inventory requires immediate monitoring.\n\n"
        "High-risk records commonly show zero on-hand inventory and zero inventory coverage.\n\n"
        "All four categories show substantial high-risk exposure.\n\n"
        "Promotional periods require stronger inventory preparation."
    ),
    (
        "Business Recommendations",
        "1. Prioritize high-risk SKUs for replenishment.\n"
        "2. Review reorder points and purchase orders.\n"
        "3. Use predicted demand for inventory planning.\n"
        "4. Monitor inventory coverage.\n"
        "5. Prepare inventory before promotions and high-risk periods."
    ),
    (
        "Streamlit Dashboard",
        "Dashboard: dashboard/app.py\n"
        "Data source: reports/risk_scoring.csv\n\n"
        "Provides demand forecasts, inventory gap, coverage, risk scores, "
        "risk levels, and inventory risk analysis.\n\n"
        "Runtime test: PASSED"
    ),
    (
        "Final Validation",
        "Final Data Validation: PASSED\n"
        "Final ML Model Validation: PASSED\n"
        "End-to-End Testing: PASSED\n"
        "Business Insights: COMPLETED\n"
        "Business Recommendations: COMPLETED\n"
        "Required project files: VERIFIED"
    ),
    (
        "Final Project Outputs",
        "models/final_random_forest.pkl\n"
        "reports/final_model_validation_metrics.csv\n"
        "reports/risk_scoring.csv\n"
        "reports/Final_Project_Report.docx\n"
        "dashboard/app.py\n"
        "README.md\n"
        "requirements.txt"
    ),
    (
        "Conclusion",
        "Project FORESIGHT provides an end-to-end machine learning solution "
        "for retail demand forecasting and inventory risk analysis, supporting "
        "demand planning, inventory monitoring, replenishment, and data-driven "
        "business decisions."
    )
]

for title, body in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    title_box = slide.shapes.add_textbox(
        Inches(0.7), Inches(0.5), Inches(12), Inches(0.8)
    )
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(30)
    title_frame.paragraphs[0].font.bold = True

    body_box = slide.shapes.add_textbox(
        Inches(0.9), Inches(1.5), Inches(11.5), Inches(5.2)
    )
    body_frame = body_box.text_frame
    body_frame.word_wrap = True
    body_frame.text = body

    for paragraph in body_frame.paragraphs:
        paragraph.font.size = Pt(20)

prs.save(output_path)

print("=" * 60)
print("FINAL PRESENTATION CREATED SUCCESSFULLY")
print("=" * 60)
print("Path:", output_path.resolve())
print("Slides:", len(prs.slides))
print("Size:", output_path.stat().st_size, "bytes")